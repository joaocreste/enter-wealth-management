"""The XP Advisory brand system, applied to a 16:9 deck.

Tokens come from `XP Advisory Brand System.html`: the dark line of the Carta ao
Investidor (a #242424 field, a copper footer band, the white symbol) for covers
and chapter openers; the light line of the Comitê de Alocação (white pages, thin
display titles in olive, hairlines instead of frames, a charcoal header band on
tables, copper as the only accent) for every page of content.
"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml
import copy

# ---------------------------------------------------------------- palette ---

INK        = RGBColor(0x22, 0x22, 0x22)
INK_BAR    = RGBColor(0x24, 0x24, 0x24)   # header / dark pages
INK_DEEP   = RGBColor(0x1F, 0x1F, 0x1F)
SLATE      = RGBColor(0x2A, 0x3B, 0x43)
COPPER     = RGBColor(0xBB, 0x79, 0x5E)
COPPER_2   = RGBColor(0xC5, 0x7D, 0x5C)
OLIVE      = RGBColor(0x82, 0x8D, 0x6F)
SAGE       = RGBColor(0xA1, 0xA8, 0x94)
CHARCOAL   = RGBColor(0x45, 0x48, 0x4A)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE   = RGBColor(0xFE, 0xFF, 0xFF)
G100       = RGBColor(0xF2, 0xF2, 0xF2)
G150       = RGBColor(0xE9, 0xE9, 0xE9)
G200       = RGBColor(0xD9, 0xD9, 0xD9)
G300       = RGBColor(0xC8, 0xC8, 0xC8)
G400       = RGBColor(0xA6, 0xA6, 0xA6)
G500       = RGBColor(0x7F, 0x7F, 0x7F)
G600       = RGBColor(0x59, 0x59, 0x59)
G700       = RGBColor(0x3B, 0x38, 0x38)
POS        = RGBColor(0x54, 0x82, 0x35)
NEG        = RGBColor(0xA6, 0x29, 0x00)
DDD        = RGBColor(0xDD, 0xDD, 0xDD)
BFB        = RGBColor(0xBF, 0xBF, 0xBF)

# ------------------------------------------------------------------ fonts ---

DISPLAY_XL = "Hanken Grotesk ExtraLight"
DISPLAY    = "Hanken Grotesk Light"
DISPLAY_R  = "Hanken Grotesk"
BODY       = "Roboto Light"
BODY_R     = "Roboto"
BODY_M     = "Roboto Medium"
MONO       = "Roboto Mono"

# ------------------------------------------------------------------- grid ---

W, H   = 13.3333, 7.5
ML     = 0.80            # left text margin
MR     = 12.24           # right edge of content
CW     = MR - ML         # content width
RAIL   = 12.60           # the vertical hairline of the light line
BAND_H = 0.26            # copper footer band on dark pages

Y_EYE   = 0.60
Y_TITLE = 0.86
Y_LEAD  = 1.64
Y_BODY  = 2.22
Y_FOOT  = 6.94


# --------------------------------------------------------------- plumbing ---

def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _noline(shape):
    shape.line.fill.background()


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        _solid(sp, fill)
    if line is None:
        _noline(sp)
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.text_frame.word_wrap = True
    sp.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    sp.text_frame.clear()
    return sp


def hairline(slide, x, y, w, color=G200, weight=0.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def vline(slide, x, y, h, color=G300, weight=0.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y + h))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def arrow(slide, x1, y1, x2, y2, color=COPPER, weight=1.0):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    lnpr = ln.line._get_or_add_ln()
    tail = lnpr.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    lnpr.append(tail)
    return ln


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.clear()
    return tb, tf


def para(tf, text="", font=BODY, size=10, color=INK, bold=False, tracking=None,
         line=1.35, before=0, after=0, align=PP_ALIGN.LEFT, first=False, caps=False):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if text:
        run(p, text, font=font, size=size, color=color, bold=bold, tracking=tracking, caps=caps)
    return p


def run(p, text, font=BODY, size=10, color=INK, bold=False, tracking=None, caps=False):
    r = p.add_run()
    r.text = text.upper() if caps else text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    if tracking:
        rPr.set('spc', str(int(round(tracking * 100))))
    # latin + east-asian + complex-script all point at the same face, so
    # PowerPoint does not substitute on a stray character
    for tag in ('a:ea', 'a:cs'):
        el = rPr.makeelement(qn(tag), {'typeface': font})
        rPr.append(el)
    return r


def fit(tf):
    """No autofit — every box on this deck is measured, not squeezed."""
    tf.auto_size = None
    return tf


# ------------------------------------------------------------ page chrome ---

def bg(slide, color):
    sp = rect(slide, 0, 0, W, H, fill=color)
    sp.shadow.inherit = False
    slide.shapes._spTree.remove(sp._element)
    slide.shapes._spTree.insert(2, sp._element)
    return sp


def light_page(slide, label, page_no, label_y=2.00):
    """White page of the Comitê line: right hairline, rotated side label, folio.

    The label reads top to bottom beside the filete, and a white knockout behind
    it interrupts the rule — the rule and the label never touch.
    """
    vline(slide, RAIL, 0.52, H - 1.04, G300, 0.75)
    text = f"{label}  —  Setembro/2026"
    extent = min(3.1, 0.070 * len(text) + 0.30)
    kn = rect(slide, RAIL - 0.13, label_y - 0.04, 0.26, extent + 0.08, fill=WHITE)
    tb, tf = textbox(slide, RAIL - extent / 2 - 0.02, label_y + extent / 2 - 0.15, extent, 0.30)
    tb.rotation = 90
    p = para(tf, "", first=True)
    run(p, label, font=BODY, size=8.5, color=OLIVE, tracking=0.6)
    run(p, "  —  Setembro/2026", font=BODY, size=8.5, color=SLATE, tracking=0.6)
    folio(slide, page_no)


def folio(slide, n):
    tb, tf = textbox(slide, RAIL - 1.0, H - 0.52, 0.86, 0.22)
    para(tf, f"{n:02d}", font=BODY_R, size=8.5, color=SLATE, align=PP_ALIGN.RIGHT, first=True)


def dark_page(slide, right_text=None, logo=None, logo_w=0.60):
    bg(slide, INK_BAR)
    band = rect(slide, 0, H - BAND_H, W, BAND_H, fill=COPPER)
    tb, tf = textbox(slide, ML, H - BAND_H, 5.0, BAND_H, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "XP ADVISORY", font=BODY, size=7.5, color=INK, tracking=2.4, first=True, caps=True)
    if right_text:
        tb2, tf2 = textbox(slide, MR - 5.0, H - BAND_H, 5.0, BAND_H, anchor=MSO_ANCHOR.MIDDLE)
        para(tf2, right_text, font=BODY_R, size=7.5, color=WHITE, tracking=2.4,
             align=PP_ALIGN.RIGHT, first=True, caps=True)
    if logo:
        slide.shapes.add_picture(logo, Inches(MR - logo_w), Inches(0.55),
                                 width=Inches(logo_w))
    return band


def eyebrow(slide, text, y=Y_EYE, x=ML, w=None, color=G500, size=8.5):
    tb, tf = textbox(slide, x, y, w or CW, 0.22)
    para(tf, text, font=BODY, size=size, color=color, tracking=size * 0.32, first=True, caps=True)
    return tb


def title(slide, text, y=Y_TITLE, x=ML, w=None, size=31, color=OLIVE, font=DISPLAY, h=0.70):
    tb, tf = textbox(slide, x, y, w or CW, h)
    para(tf, text, font=font, size=size, color=color, line=1.06, first=True)
    return tb


def lead(slide, text, y=Y_LEAD, x=ML, w=None, size=11.5, color=INK, h=0.52, font=BODY):
    tb, tf = textbox(slide, x, y, w or CW, h)
    para(tf, text, font=font, size=size, color=color, line=1.45, first=True)
    return tb


def source(slide, text, y=Y_FOOT, x=ML, w=None):
    tb, tf = textbox(slide, x, y, w or CW, 0.20)
    p = para(tf, "", first=True)
    run(p, "Fonte: ", font=BODY_R, size=7.5, color=G600)
    run(p, text, font=BODY, size=7.5, color=G600)
    return tb


# ------------------------------------------------------------- components ---

def callout(slide, label, body, x, y, w, h, lab_w=1.72, lab_fill=SLATE,
            body_size=9.2, lab_size=7.5, rule=True, body_font=BODY, body_color=INK):
    """The GAAC callout: a slate label block, a body, a hairline under both."""
    box = rect(slide, x, y, lab_w, h, fill=lab_fill)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = Inches(0.07)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, label, font=BODY_R, size=lab_size, color=OFFWHITE, bold=True,
         tracking=lab_size * 0.16, line=1.25, align=PP_ALIGN.CENTER, first=True, caps=True)
    tb, tf2 = textbox(slide, x + lab_w + 0.22, y + 0.035, w - lab_w - 0.22, h - 0.07,
                      anchor=MSO_ANCHOR.MIDDLE)
    for i, line_text in enumerate(body if isinstance(body, (list, tuple)) else [body]):
        para(tf2, line_text, font=body_font, size=body_size, color=body_color,
             line=1.42, first=(i == 0), before=0 if i == 0 else 3)
    if rule:
        hairline(slide, x, y + h, w, G200, 0.75)
    return box


def numbered(slide, number, head, body, x, y, w, h, num_color=COPPER, head_size=12.5,
             body_size=9.2, rule=True, head_color=INK, num_size=9):
    tb, tf = textbox(slide, x, y, 0.60, 0.24)
    para(tf, number, font=MONO, size=num_size, color=num_color, tracking=num_size * 0.12, first=True)
    tb2, tf2 = textbox(slide, x + 0.60, y - 0.045, w - 0.60, 0.30)
    para(tf2, head, font=DISPLAY_R, size=head_size, color=head_color, line=1.1, first=True)
    tb3, tf3 = textbox(slide, x + 0.60, y + 0.30, w - 0.60, h - 0.34)
    for i, line_text in enumerate(body if isinstance(body, (list, tuple)) else [body]):
        para(tf3, line_text, font=BODY, size=body_size, color=INK, line=1.45,
             first=(i == 0), before=0 if i == 0 else 4)
    if rule:
        hairline(slide, x, y + h, w, G200, 0.75)


def kpi(slide, figure, caption, x, y, w, fig_size=30, fig_color=COPPER):
    hairline(slide, x, y, w, G300, 0.75)
    tb, tf = textbox(slide, x, y + 0.14, w, 0.52)
    para(tf, figure, font=DISPLAY, size=fig_size, color=fig_color, line=1.0, first=True)
    tb2, tf2 = textbox(slide, x, y + 0.70, w, 0.50)
    para(tf2, caption, font=BODY, size=8.6, color=G600, line=1.35, first=True)


def chip(slide, text, x, y, w, h=0.30, fill=G150, color=INK_BAR, size=8.5, font=DISPLAY_R,
         tracking=None, bold=False):
    sp = rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    sp.adjustments[0] = 0.5
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, font=font, size=size, color=color, align=PP_ALIGN.CENTER,
         tracking=tracking, bold=bold, line=1.2, first=True)
    return sp


def band(slide, text, x, y, w, h=0.34, fill=G100, color=G600, size=8, tracking=None):
    sp = rect(slide, x, y, w, h, fill=fill)
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.10)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, font=BODY, size=size, color=color, align=PP_ALIGN.CENTER,
         tracking=tracking if tracking is not None else size * 0.28, line=1.2,
         first=True, caps=True)
    return sp


def card(slide, x, y, w, h, fill=WHITE, line=G200):
    return rect(slide, x, y, w, h, fill=fill, line=line, line_w=0.75)


def picture(slide, path, x, y, w=None, h=None, frame=True):
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y),
                                   width=Inches(w) if w else None,
                                   height=Inches(h) if h else None)
    if frame:
        ln = pic.line
        ln.color.rgb = G200
        ln.width = Pt(0.75)
    return pic


def caption(slide, text, x, y, w, align=PP_ALIGN.LEFT):
    tb, tf = textbox(slide, x, y, w, 0.20)
    para(tf, text, font=BODY, size=7.5, color=G500, tracking=1.9, align=align,
         first=True, caps=True)
    return tb


# ------------------------------------------------------------- measuring ---

def est_lines(text, width_in, size_pt, factor=0.485):
    """Lines a string needs in a column, from the face's average advance."""
    per_line = max(8, int(width_in * 72 / (size_pt * factor)))
    return max(1, -(-len(text) // per_line))


def text_h(text, width_in, size_pt, line=1.45):
    return est_lines(text, width_in, size_pt) * size_pt * line / 72


def bullets(slide, items, x, y, w, size=9.8, color=COPPER, line=1.45, gap=0.20,
            indent=0.24, mark=0.075):
    """Copper square, text, and as much room as the text asks for."""
    for it in items:
        h = text_h(it, w - indent, size, line)
        rect(slide, x + 0.02, y + (size * line / 72) / 2 - mark / 2, mark, mark, fill=color)
        tb, tf = textbox(slide, x + indent, y, w - indent, h + 0.06)
        para(tf, it, font=BODY, size=size, color=INK, line=line, first=True)
        y += h + gap
    return y


# ----------------------------------------------------------------- icons ---
#
# The icons are drawn, not placed. An SVG path on a 24-unit grid is converted
# to a PowerPoint freeform (a:custGeom), so a mark is a vector shape like any
# other on the page: it survives a zoom, it recolours with the palette, and it
# carries the same hairline weight as the rules. A PNG would have been quicker
# and would have been the one pixelated thing on a deck that is otherwise all
# vector.
#
# Stroke only — no icon is filled — which is what keeps them in the register of
# the light line: thin olive and copper on white, never a solid block.

import re

ICON_GRID = 24        # the viewBox every path below is written on
_PATH_SPACE = 21600   # the units PowerPoint scales to the shape's extent
_TOKEN = re.compile(r'[A-Za-z]|-?\d*\.?\d+')
_K = 0.5523           # circular arc as a cubic bézier


def _circle(cx, cy, r):
    k = r * _K
    return (f"M {cx},{cy - r} "
            f"C {cx + k},{cy - r} {cx + r},{cy - k} {cx + r},{cy} "
            f"C {cx + r},{cy + k} {cx + k},{cy + r} {cx},{cy + r} "
            f"C {cx - k},{cy + r} {cx - r},{cy + k} {cx - r},{cy} "
            f"C {cx - r},{cy - k} {cx - k},{cy - r} {cx},{cy - r} Z")


def _ellipse(cx, cy, rx, ry):
    kx, ky = rx * _K, ry * _K
    return (f"M {cx},{cy - ry} "
            f"C {cx + kx},{cy - ry} {cx + rx},{cy - ky} {cx + rx},{cy} "
            f"C {cx + rx},{cy + ky} {cx + kx},{cy + ry} {cx},{cy + ry} "
            f"C {cx - kx},{cy + ry} {cx - rx},{cy + ky} {cx - rx},{cy} "
            f"C {cx - rx},{cy - ky} {cx - kx},{cy - ry} {cx},{cy - ry} Z")


def _rrect(x, y, w, h, r):
    k = r * (1 - _K)
    return (f"M {x + r},{y} L {x + w - r},{y} "
            f"C {x + w - k},{y} {x + w},{y + k} {x + w},{y + r} "
            f"L {x + w},{y + h - r} "
            f"C {x + w},{y + h - k} {x + w - k},{y + h} {x + w - r},{y + h} "
            f"L {x + r},{y + h} "
            f"C {x + k},{y + h} {x},{y + h - k} {x},{y + h - r} "
            f"L {x},{y + r} "
            f"C {x},{y + k} {x + k},{y} {x + r},{y} Z")


# Each icon says what the thing is, in the plainest mark that says it: the
# cylinder for the relational store, the archive box for the bucket of finished
# artefacts, the key for the key-value cache, the asterisk for the model.
ICONS = {
    # a clock — the cron that starts the day
    'clock': _circle(12, 12, 8.6) + " M 12,6.6 L 12,12.2 L 16,14.4",

    # a pointer — the advisor, who asks for the letter from the portal
    'cursor': "M 5.4,2.8 L 5.4,18.4 L 9.6,14.3 L 12.3,20.2 L 14.9,19 L 12.2,13.2 L 18,12.9 Z",

    # a terminal — the Node job that deposits what the Worker is refused
    'terminal': _rrect(2.4, 4.4, 19.2, 15.2, 2.2) + " M 7,10.2 L 10.4,12.9 L 7,15.6 M 13,15.6 L 17.4,15.6",

    # a chip — the Worker, where the account is done
    'chip': (_rrect(6.4, 6.4, 11.2, 11.2, 1.6) + " " + _rrect(10.2, 10.2, 3.6, 3.6, 0.5) +
             " M 9.4,6.4 L 9.4,3.2 M 14.6,6.4 L 14.6,3.2 M 9.4,20.8 L 9.4,17.6 M 14.6,20.8 L 14.6,17.6"
             " M 6.4,9.4 L 3.2,9.4 M 6.4,14.6 L 3.2,14.6 M 20.8,9.4 L 17.6,9.4 M 20.8,14.6 L 17.6,14.6"),

    # three linked nodes — a Workflow, one durable step after another
    'flow': (_circle(4.6, 12, 2.4) + " " + _circle(12, 12, 2.4) + " " + _circle(19.4, 12, 2.4) +
             " M 7.2,12 L 9.4,12 M 14.6,12 L 16.8,12"),

    # the cylinder — D1, the relational record
    'database': (_ellipse(12, 6.2, 8, 3.2) +
                 " M 4,6.2 L 4,17.8 C 4,19.57 7.58,21 12,21 C 16.42,21 20,19.57 20,17.8 L 20,6.2"
                 " M 20,12 C 20,13.77 16.42,15.2 12,15.2 C 7.58,15.2 4,13.77 4,12"),

    # the archive box — R2, where a finished artefact is filed and not touched again
    'bucket': (_rrect(2.4, 3.4, 19.2, 4.4, 1.0) +
               " M 4.4,7.8 L 4.4,19.2 C 4.4,20.19 5.21,21 6.2,21 L 17.8,21"
               " C 18.79,21 19.6,20.19 19.6,19.2 L 19.6,7.8 M 9.8,11.8 L 14.2,11.8"),

    # the key — KV, fetched by its name and never authoritative
    'key': (_circle(7.6, 12, 3.8) + " M 11.4,12 L 21,12 M 17.6,12 L 17.6,15.6 M 20.2,12 L 20.2,14.8"),

    # the asterisk — the model, which writes and classifies and nothing else
    'spark': "M 12,2.8 L 12,21.2 M 4.04,7.4 L 19.96,16.6 M 4.04,16.6 L 19.96,7.4",

    # the globe — the providers outside, each one named beside its number
    'globe': (_circle(12, 12, 8.8) + " " + _ellipse(12, 12, 3.9, 8.8) +
              " M 3.6,9 L 20.4,9 M 3.6,15 L 20.4,15"),
}


def _path_xml(d):
    """SVG path data (absolute M, L, H, V, C and Z) as DrawingML path commands."""
    u = lambda v: int(round(float(v) * _PATH_SPACE / ICON_GRID))
    pt = lambda x, y: f'<a:pt x="{u(x)}" y="{u(y)}"/>'
    toks, i, cmd, out = _TOKEN.findall(d), 0, None, []
    x = y = 0.0
    while i < len(toks):
        if toks[i].isalpha():
            cmd, i = toks[i], i + 1
            if cmd == 'Z':
                out.append('<a:close/>')
                continue
        n = {'M': 2, 'L': 2, 'H': 1, 'V': 1, 'C': 6}.get(cmd)
        if n is None:
            raise ValueError(f"icon path: unsupported command {cmd!r} — write it as M, L, H, V, C or Z")
        a = [float(v) for v in toks[i:i + n]]
        i += n
        if cmd == 'M':
            x, y = a
            out.append(f'<a:moveTo>{pt(x, y)}</a:moveTo>')
        elif cmd == 'L':
            x, y = a
            out.append(f'<a:lnTo>{pt(x, y)}</a:lnTo>')
        elif cmd == 'H':
            x = a[0]
            out.append(f'<a:lnTo>{pt(x, y)}</a:lnTo>')
        elif cmd == 'V':
            y = a[0]
            out.append(f'<a:lnTo>{pt(x, y)}</a:lnTo>')
        else:
            out.append(f'<a:cubicBezTo>{pt(a[0], a[1])}{pt(a[2], a[3])}{pt(a[4], a[5])}</a:cubicBezTo>')
            x, y = a[4], a[5]
    return ''.join(out)


def icon(slide, name, x, y, size, color=COPPER, weight=1.0):
    """Draw an icon at (x, y), `size` inches square, stroked in `color`."""
    if name not in ICONS:
        raise KeyError(f"no icon named {name!r} — have {', '.join(sorted(ICONS))}")
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(size), Inches(size))
    sp.shadow.inherit = False
    spPr = sp._element.spPr
    prst = spPr.find(qn('a:prstGeom'))
    geom = parse_xml(
        f'<a:custGeom {nsdecls("a")}><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
        f'<a:rect l="0" t="0" r="r" b="b"/><a:pathLst>'
        f'<a:path w="{_PATH_SPACE}" h="{_PATH_SPACE}" fill="none" stroke="1">{_path_xml(ICONS[name])}</a:path>'
        f'</a:pathLst></a:custGeom>')
    spPr.insert(list(spPr).index(prst), geom)
    spPr.remove(prst)
    sp.fill.background()
    sp.line.color.rgb = color
    sp.line.width = Pt(weight)
    ln = sp.line._get_or_add_ln()
    ln.set('cap', 'rnd')                                   # the stroke ends soft, like the rules
    ln.append(ln.makeelement(qn('a:round'), {}))
    sp.text_frame.clear()
    return sp
