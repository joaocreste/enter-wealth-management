"""Builds the XP deck: what is happening, what we recommend, how it works, next steps."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from brand import *  # noqa: F401,F403
import brand as B

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
OUT = os.environ.get("DECK_OUT", os.path.join(HERE, "deck.pptx"))

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

page = {"n": 0}


def new(dark=False, label=None, right=None, logo=True, label_y=2.00):
    s = prs.slides.add_slide(BLANK)
    page["n"] += 1
    if dark:
        dark_page(s, right_text=right,
                  logo=os.path.join(IMG, "xp-logo-white.png") if logo else None)
    else:
        light_page(s, label, page["n"], label_y=label_y)
    return s


def block(slide, num, head, items, x, y, w, head_size=16, size=9.8, rule=True):
    """A numbered front: figure, display head, measured bullets, closing rule."""
    tb, tf = textbox(slide, x, y + 0.03, 0.60, 0.26)
    para(tf, num, font=MONO, size=10, color=COPPER, tracking=1.2, first=True)
    tb, tf = textbox(slide, x + 0.62, y - 0.05, w - 0.62, 0.34)
    para(tf, head, font=DISPLAY_R, size=head_size, color=INK, line=1.08, first=True)
    end = bullets(slide, items, x + 0.64, y + 0.46, w - 0.64, size=size)
    if rule:
        hairline(slide, x, end + 0.10, w, G200, 0.75)
    return end + 0.34


# ═══════════════════════════════════════════════════════════ 01 · capa ═══

s = new(dark=True, right="15 de setembro de 2026")
eyebrow(s, "XP Asset Management  ·  Wealth Management", y=1.42, color=SAGE, size=9)
tb, tf = textbox(s, ML, 1.86, 10.2, 2.0)
para(tf, "A carta mensal,", font=DISPLAY_XL, size=54, color=WHITE, line=1.06, first=True)
para(tf, "para o book inteiro", font=DISPLAY_XL, size=54, color=WHITE, line=1.06)
hairline(s, ML, 4.22, 1.30, COPPER, 1.5)
tb, tf = textbox(s, ML, 4.52, 8.6, 1.4)
para(tf, "Um produto para o assessor escrever, todo mês, a cada um de até trezentos "
         "clientes — com a aritmética feita em código, a fonte de cada número anotada "
         "e a aprovação do assessor antes de qualquer envio.",
     font=BODY, size=13, color=BFB, line=1.55, first=True)
tb, tf = textbox(s, ML, 6.10, 8.6, 0.30)
para(tf, "Sessão de trabalho Enter × XP  ·  construído em Claude Code, sobre a API da Anthropic",
     font=BODY, size=8.5, color=G600, tracking=1.7, first=True, caps=True)


# ══════════════════════════════════════════════════════ 02 · o roteiro ═══

s = new(label="Roteiro")
eyebrow(s, "Roteiro  ·  quatro perguntas")
title(s, "O que vamos percorrer")
lead(s, "Onde estamos, o que recomendamos, como o que foi construído funciona, e o que "
        "separa o que está de pé do primeiro cliente real.", h=0.40)

ROT = [
    ("01", "What is happening and why are we here?",
     "Trezentas cartas por mês, por assessor. A primeira tentativa rodava, produzia "
     "prosa em português e não podia ser mostrada a um cliente."),
    ("02", "What do you recommend?",
     "Um produto novo, formado por duas pontas que se encontraram no meio: o mercado "
     "lido todo dia, a carteira de cada cliente, e a carta entre os dois."),
    ("03", "How it works",
     "Quatro agentes. Três montam o panorama do dia sozinhos na nuvem; o quarto "
     "consolida a carta mensal quando o assessor manda."),
    ("04", "Next steps",
     "O que falta para sair do dado simulado: acesso, contrato de nuvem, segurança "
     "do dado financeiro, sigilo e governança do que o modelo escreve."),
]
y = 2.46
for num, head, body in ROT:
    numbered(s, num, head, body, ML, y, CW, 1.02, head_size=16, body_size=10)
    y += 1.14


# ═══════════════════════════════════════════════ 03 · abertura · I ═══════

s = new(dark=True, right="Capítulo 01")
eyebrow(s, "Capítulo 01", y=2.30, color=COPPER, size=9)
tb, tf = textbox(s, ML, 2.66, 10.6, 1.5)
para(tf, "What is happening", font=DISPLAY_XL, size=42, color=WHITE, line=1.1, first=True)
para(tf, "and why are we here?", font=DISPLAY_XL, size=42, color=WHITE, line=1.1)
hairline(s, ML, 4.44, 1.30, COPPER, 1.5)
tb, tf = textbox(s, ML, 4.72, 8.8, 0.8)
para(tf, "Trezentas cartas por mês, por assessor — e uma primeira tentativa que não "
         "sobreviveu à leitura.", font=BODY, size=12.5, color=BFB, line=1.5, first=True)


# ══════════════════════════════════════════════ 04 · a escala do problema ═

s = new(label="Contexto")
eyebrow(s, "O que está acontecendo  ·  a escala")
title(s, "Trezentas cartas por mês, por assessor")
lead(s, "Cada assessor mantém um book de até trezentos clientes, e cada cliente espera "
        "uma carta por mês. A conta não fecha à mão, e o que a carta precisa conter não é "
        "opinião: é aritmética, procedência e política.")

KPIS = [
    ("300", "clientes no book de um assessor"),
    ("12", "cartas por cliente, por ano"),
    ("3.600", "cartas por assessor, por ano"),
    ("14", "cartas por dia útil, se escritas à mão"),
]
gap = 0.30
kw = (CW - gap * 3) / 4
for i, (fig, cap) in enumerate(KPIS):
    kpi(s, fig, cap, ML + i * (kw + gap), 2.52, kw, fig_size=34)

hairline(s, ML, 4.14, CW, G300, 0.75)
tb, tf = textbox(s, ML, 4.32, CW, 0.30)
para(tf, "E cada uma dessas cartas precisa acertar quatro coisas", font=BODY_R,
     size=9, color=G600, tracking=1.8, first=True, caps=True)

NEEDS = [
    ("Rentabilidade", "O retorno do mês da carteira, medido contra a referência da "
                      "política que o próprio cliente assinou."),
    ("Atribuição", "Quem puxou o resultado para cima e para baixo, posição a posição, "
                   "com o efeito do câmbio separado."),
    ("Mercado", "O que aconteceu no mês e, de tudo aquilo, o que de fato toca esta "
                "carteira — e não a de todo mundo."),
    ("Proposta", "O que sugerir, dentro da faixa, do teto e das restrições, com o "
                 "sinal de mercado e o consenso ao lado."),
]
y = 4.74
for lab, body in NEEDS:
    callout(s, lab, body, ML, y, CW, 0.52, lab_w=1.55, body_size=9.2)
    y += 0.56

source(s, "volume declarado pela XP na sessão de trabalho; dias úteis, média de um mês comercial.")


# ═════════════════════════════════════════════ 05 · a primeira tentativa ══

s = new(label="Contexto")
eyebrow(s, "O que está acontecendo  ·  a primeira tentativa")
title(s, "Rodava, e não podia ser mostrada a um cliente")
lead(s, "Quatro prompts em cadeia: resuma o perfil, resuma o macro, resuma a carteira, "
        "escreva a carta. Produzia prosa em português e falhava em tudo o que uma carta "
        "precisa acertar.", h=0.46)

FAILS = [
    ("Sem aritmética", "Nenhum retorno era calculado. O extrato traz rentabilidade desde o "
                       "início — LREN3 a −41,7%, HAPV3 a −74,58% — e nada impedia que "
                       "aquilo fosse apresentado como o resultado do mês."),
    ("Sem fontes", "Cada número era o que o modelo repetia da entrada, sem registro de "
                   "origem, de quando foi buscado ou se ainda era verdade. O relatório "
                   "macro era de fevereiro de 2025 e entrava como se fosse de hoje."),
    ("Dado vencido", "MRFG3 virou MBRF3 e ARZZ3 virou AZZA3; nenhum provedor devolve preço "
                     "para os tickers antigos. O CDB do extrato já havia vencido. A "
                     "avaliação de maio de 2025 seguiria adiante, em silêncio."),
    ("Sem política", "“Faça recomendações alinhadas ao perfil” deixava o modelo inventá-las: "
                     "sem sinal de mercado, sem faixa da política, sem checagem de "
                     "concentração e sem lista de restrições."),
    ("Sem formato", "A saída era texto de chat. As duas páginas, o formato da carta e a "
                    "marca eram trabalho manual, refeito a cada cliente."),
    ("Sem estado", "Nada era versionado, nada era guardado, nada era reproduzível — e não "
                   "havia assessor no circuito: o primeiro rascunho do modelo era a entrega."),
]
y = 2.40
for lab, body in FAILS:
    callout(s, lab, body, ML, y, CW, 0.66, lab_w=1.55, body_size=9.2)
    y += 0.68

tb, tf = textbox(s, ML, y + 0.18, CW, 0.40)
p = para(tf, "", first=True)
run(p, "Debaixo de tudo, um erro de projeto:  ", font=BODY, size=12.5, color=INK)
run(p, "o modelo era o sistema.", font=DISPLAY_R, size=13.5, color=COPPER)
source(s, "docs/implementation-report.md §1 — revisão da versão inicial.")


# ═══════════════════════════════════════════════ 06 · abertura · II ══════

s = new(dark=True, right="Capítulo 02")
eyebrow(s, "Capítulo 02", y=2.30, color=COPPER, size=9)
tb, tf = textbox(s, ML, 2.66, 10.6, 1.5)
para(tf, "What do you", font=DISPLAY_XL, size=42, color=WHITE, line=1.1, first=True)
para(tf, "recommend?", font=DISPLAY_XL, size=42, color=WHITE, line=1.1)
hairline(s, ML, 4.44, 1.30, COPPER, 1.5)
tb, tf = textbox(s, ML, 4.72, 8.8, 0.8)
para(tf, "Um produto novo, formado por duas pontas que se encontraram no meio.",
     font=BODY, size=12.5, color=BFB, line=1.5, first=True)


# ═════════════════════════════════════════ 07 · duas pontas, uma carta ═══

s = new(label="Recomendação")
eyebrow(s, "A recomendação  ·  a forma do produto")
title(s, "Duas pontas, e a carta no meio")
lead(s, "De um lado, o mercado, lido todo dia. Do outro, a carteira de cada cliente e a "
        "política que a governa. O produto é o encontro dos dois, e o encontro tem a "
        "forma de uma carta — que o assessor lê, corrige e aprova antes de o cliente ver.",
     h=0.46)

BOXY, BOXH = 2.56, 3.14
LW, CW2 = 3.34, 3.72
LX = ML
CX = ML + LW + 0.34
RX = CX + CW2 + 0.34


def side_box(x, head, items):
    card(s, x, BOXY + 0.24, LW, BOXH - 0.24)
    hd = rect(s, x, BOXY + 0.24, LW, 0.42, fill=SLATE)
    tf = hd.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, head, font=BODY_R, size=8.5, color=OFFWHITE, bold=True, tracking=1.7,
         align=PP_ALIGN.CENTER, first=True, caps=True)
    tb, tf2 = textbox(s, x + 0.24, BOXY + 0.86, LW - 0.48, BOXH - 1.12)
    for i, it in enumerate(items):
        para(tf2, it, font=BODY, size=9.4, color=INK, line=1.42, first=(i == 0),
             before=0 if i == 0 else 10)


side_box(LX, "O mercado", [
    "Cotações, séries diárias e câmbio, de provedor nomeado.",
    "Leitura técnica e consenso de analistas, por ativo.",
    "As manchetes das últimas 48 horas, cada uma ligada ao seu artigo.",
    "Indicadores medidos em 5 sessões e em 30 dias.",
])
side_box(RX, "O cliente", [
    "A carteira aprovada, posição a posição, com os fluxos do mês.",
    "A política de investimento: faixas, tetos e restrições.",
    "O perfil de suitability e a última reunião registrada.",
    "A exposição por classe de ativo, hoje.",
])

ctr = rect(s, CX, BOXY, CW2, BOXH, fill=INK_BAR)
rect(s, CX, BOXY, CW2, 0.065, fill=COPPER)
tb, tf = textbox(s, CX + 0.28, BOXY + 0.44, CW2 - 0.56, 0.36)
para(tf, "A carta mensal", font=DISPLAY, size=19, color=WHITE, line=1.1,
     align=PP_ALIGN.CENTER, first=True)
hairline(s, CX + 1.30, BOXY + 0.98, CW2 - 2.60, COPPER, 1.0)
tb, tf = textbox(s, CX + 0.30, BOXY + 1.16, CW2 - 0.60, 1.80)
for i, it in enumerate([
    "Rentabilidade e atribuição do mês, contra a referência da própria política.",
    "O que do mercado toca esta carteira — e só isso.",
    "Propostas medidas contra a faixa, o teto e as restrições.",
    "Lida, corrigida e aprovada pelo assessor antes de sair.",
]):
    para(tf, it, font=BODY, size=9.4, color=DDD, line=1.4, align=PP_ALIGN.CENTER,
         first=(i == 0), before=0 if i == 0 else 9)

arrow(s, LX + LW + 0.06, BOXY + BOXH / 2, CX - 0.06, BOXY + BOXH / 2, COPPER, 1.25)
arrow(s, RX - 0.06, BOXY + BOXH / 2, CX + CW2 + 0.06, BOXY + BOXH / 2, COPPER, 1.25)

hairline(s, ML, 5.98, CW, G200, 0.75)
tb, tf = textbox(s, ML, 6.16, CW, 0.80)
p = para(tf, "", first=True)
run(p, "Duas portas para o mesmo documento. ", font=BODY_R, size=10.5, color=INK)
run(p, "No portal do assessor a carta é feita, revisada e aprovada; no portal do cliente "
       "ela é lida, junto da carteira e do mês. Um documento por cliente por mês, um "
       "lugar só onde ele existe — e o mesmo agente percorre o book inteiro quando o "
       "assessor pede todas as cartas de uma vez.",
    font=BODY, size=10.5, color=INK)


# ═══════════════════════════════════════ 08 · look & feel · assessor ═════

TW = 7.10           # the text column of the two look-and-feel pages
SHOT_X, SHOT_W = 8.94, 3.30
SHOT_Y, SHOT_H = 2.44, 4.46

s = new(label="Recomendação")
eyebrow(s, "O que foi criado  ·  portal do assessor")
title(s, "Portal do assessor")
lead(s, "Trilho escuro à esquerda com o símbolo branco, páginas claras, títulos finos em "
        "oliva, filetes no lugar de molduras, tabelas com faixa carvão e o cobre como "
        "único acento.", w=TW, h=0.62)

PAGES_A = [
    ("Panorama do dia", "O resumo do dia escrito pelo agente de inferência; O que importa "
                        "hoje em duas tabelas, Brasil e internacional; os indicadores com "
                        "sparkline; os gatilhos em barras; os desvios por cliente e a "
                        "matriz de correlações."),
    ("Sinais de mercado", "Leitura técnica e consenso de analistas por ativo, duas famílias "
                          "independentes que podem discordar, e o risco-retorno dos doze "
                          "meses fechados com a fronteira empírica por cima."),
    ("Clientes", "A tabela do book — patrimônio, mês, referência, última reunião, próxima "
                 "revisão, estado da carta e alertas — e um cartão por cliente com os "
                 "pontos que pedem atenção."),
    ("Gatilhos", "Cada limiar configurado e cada desvio de alocação com sua proximidade "
                 "(1,0 = no limiar), desenhado como barra, dizendo quando uma ação é devida."),
    ("Página do cliente", "Carteira, desempenho, recomendações com o guarda-corpo de "
                          "suitability, preparação de reunião, editor de carteira, cartas e "
                          "auditoria — e os dois botões que disparam agentes."),
]
y = 2.50
for lab, body in PAGES_A:
    callout(s, lab, body, ML, y, TW, 0.80, lab_w=1.42, body_size=8.8)
    y += 0.84

caption(s, "Panorama do dia  ·  topo da página", SHOT_X, 2.18, SHOT_W)
picture(s, os.path.join(IMG, "shot-panorama.png"), SHOT_X, SHOT_Y, h=SHOT_H)


# ════════════════════════════════════════ 09 · look & feel · cliente ═════

s = new(label="Recomendação")
eyebrow(s, "O que foi criado  ·  portal do cliente e a carta")
title(s, "Portal do cliente, e a carta")
lead(s, "O mesmo sistema visual com menos superfície. O cliente vê a carteira, o mês e a "
        "carta que o assessor aprovou — nada em rascunho, nada gerado na hora em que a "
        "página abre.", w=TW, h=0.62)

PAGES_C = [
    ("Minha carteira", "A composição por classe com a faixa permitida desenhada em vez de "
                       "escrita, as posições e o patrimônio do dia, cada preço com o "
                       "provedor que o forneceu."),
    ("Último mês", "Rentabilidade contra a referência da própria política, atribuição por "
                   "posição e os fluxos do mês — os mesmos números da carta, porque saem "
                   "do mesmo cálculo."),
    ("O que importa", "O recorte do panorama do dia que toca esta carteira, com a fonte de "
                      "cada linha e o link para a matéria de onde a manchete veio."),
    ("Carta do assessor", "A carta publicada, exatamente como foi aprovada, com o gráfico "
                          "de posicionamento, as fontes usadas ao pé e o disclaimer."),
    ("Documentos", "A política de investimento em vigor — o PDF em que ela foi acordada, "
                   "quem subiu e quando — e as cartas anteriores."),
]
y = 2.50
for lab, body in PAGES_C:
    callout(s, lab, body, ML, y, TW, 0.74, lab_w=1.42, body_size=8.8)
    y += 0.78

tb, tf = textbox(s, ML, y + 0.18, TW, 0.80)
p = para(tf, "", first=True)
run(p, "Uma carga, três saídas. ", font=BODY_R, size=9.8, color=INK)
run(p, "Portal, PDF de duas páginas e e-mail saem do mesmo modelo de carta e da mesma "
       "geometria de gráfico, de modo que não podem discordar entre si.",
    font=BODY, size=9.8, color=INK)

caption(s, "Carta do assessor  ·  e o PDF", SHOT_X, 2.18, SHOT_W)
picture(s, os.path.join(IMG, "shot-letter.png"), SHOT_X, SHOT_Y, h=SHOT_H)
# the PDF sits over the letter's lower corner, on a white edge that separates them
rect(s, 10.80, 4.78, 1.40, 1.93, fill=WHITE)
picture(s, os.path.join(IMG, "pdf1.png"), 10.86, 4.84, w=1.28)


# ═══════════════════════════════════════════════ 10 · abertura · III ═════

s = new(dark=True, right="Capítulo 03")
eyebrow(s, "Capítulo 03", y=2.30, color=COPPER, size=9)
tb, tf = textbox(s, ML, 2.66, 10.6, 1.5)
para(tf, "How it works", font=DISPLAY_XL, size=42, color=WHITE, line=1.1, first=True)
hairline(s, ML, 3.86, 1.30, COPPER, 1.5)
tb, tf = textbox(s, ML, 4.14, 8.8, 0.8)
para(tf, "Quatro agentes: três montam o panorama do dia e rodam sozinhos na nuvem; o "
         "quarto consolida a carta mensal quando o assessor manda.",
     font=BODY, size=12.5, color=BFB, line=1.5, first=True)


# ═══════════════════════════════════════════ 11 · os quatro agentes ══════

s = new(label="Como funciona")
eyebrow(s, "Como funciona  ·  o esquema")
title(s, "Quatro agentes")

band(s, "Todo dia às 07:00 de São Paulo  ·  cron da Cloudflare  ·  e de novo sempre que o "
        "assessor pedir", ML, 1.72, CW, 0.32, fill=G100, color=G600, size=8)

AG = [
    ("01", "Dados", "Busca cotações, séries diárias e sinais técnicos, mede 5 sessões e 30 "
                    "dias, lê as manchetes das últimas 48 horas — e anota a fonte de cada "
                    "número que traz."),
    ("02", "Inferência", "Lê o que o primeiro reuniu junto da exposição de cada cliente e "
                         "decide o que, de tudo aquilo, importa para estas carteiras hoje. "
                         "É ele que escreve o resumo do dia."),
    ("03", "Gatilhos", "Compara cada indicador com os limiares configurados de cada carteira "
                       "e cada desvio com a política que a governa, e diz quando uma ação é "
                       "devida."),
]
bw = (CW - 0.46 * 2) / 3
for i, (num, head, body) in enumerate(AG):
    x = ML + i * (bw + 0.46)
    card(s, x, 2.28, bw, 1.86)
    rect(s, x, 2.28, bw, 0.05, fill=COPPER if i < 2 else SLATE)
    tb, tf = textbox(s, x + 0.24, 2.50, bw - 0.48, 0.24)
    para(tf, f"agente {num}", font=MONO, size=8, color=COPPER, tracking=1.0, first=True, caps=True)
    tb, tf = textbox(s, x + 0.24, 2.76, bw - 0.48, 0.32)
    para(tf, head, font=DISPLAY_R, size=17, color=INK, line=1.05, first=True)
    tb, tf = textbox(s, x + 0.24, 3.20, bw - 0.48, 0.82)
    para(tf, body, font=BODY, size=9, color=INK, line=1.42, first=True)
    if i < 2:
        arrow(s, x + bw + 0.06, 3.21, x + bw + 0.40, 3.21, COPPER, 1.25)

hairline(s, ML, 4.32, CW, G300, 0.75)
tb, tf = textbox(s, ML, 4.42, CW, 0.24)
p = para(tf, "", first=True)
run(p, "Panorama do dia", font=BODY_R, size=8.5, color=SLATE, tracking=1.7, caps=True)
run(p, "   ·   rodam em sequência, sozinhos, sem depender da máquina de ninguém",
    font=BODY, size=8.5, color=G600, tracking=1.7, caps=True)

Y4 = 5.00
card(s, ML, Y4, CW, 1.54, fill=G100, line=G200)
rect(s, ML, Y4, 0.055, 1.54, fill=COPPER)
tb, tf = textbox(s, ML + 0.30, Y4 + 0.22, 4.4, 0.24)
para(tf, "agente 04", font=MONO, size=8, color=COPPER, tracking=1.0, first=True, caps=True)
tb, tf = textbox(s, ML + 0.30, Y4 + 0.46, 4.6, 0.34)
para(tf, "Carta mensal", font=DISPLAY_R, size=17, color=INK, line=1.05, first=True)
tb, tf = textbox(s, ML + 0.30, Y4 + 0.88, 4.5, 0.52)
para(tf, "Comandado pelo assessor a partir da página do cliente. Abre uma aba própria e "
         "relata cada etapa enquanto roda.", font=BODY, size=8.8, color=INK, line=1.4, first=True)

STEPS = ["Dados", "Análise", "Redação", "Diagramação"]
sx, sw, sg = 6.10, 1.32, 0.26
for i, st in enumerate(STEPS):
    x = sx + i * (sw + sg)
    chip(s, st, x, Y4 + 0.44, sw, 0.36, fill=WHITE, color=INK, size=9.5, font=DISPLAY_R)
    if i < 3:
        arrow(s, x + sw + 0.04, Y4 + 0.62, x + sw + sg - 0.04, Y4 + 0.62, COPPER, 1.0)
tb, tf = textbox(s, sx, Y4 + 0.92, 12.16 - sx, 0.50)
para(tf, "Reúne os dados do mês, calcula rentabilidade, atribuição e propostas, escreve a "
         "carta, diagrama em PDF e e-mail e atualiza o portal. É esse mesmo agente que "
         "percorre o book inteiro quando o assessor pede todas as cartas de uma vez.",
     font=BODY, size=8.8, color=INK, line=1.4, first=True)


# ═══════════════════════════════════════════ 12 · os três do dia ═════════

s = new(label="Como funciona")
eyebrow(s, "Como funciona  ·  os três agentes do dia")
title(s, "O panorama, montado sozinho")
lead(s, "Rodam em sequência dentro do Worker, todo dia às sete da manhã e de novo sempre "
        "que o assessor aperta atualizar dados de mercado — o que abre um cartão de "
        "progresso dizendo o que cada agente está fazendo enquanto roda.", h=0.46)

DETAIL = [
    ("01", "Dados", COPPER, [
        ("O que busca", "Todo indicador monitorado com o movimento do dia, as séries "
                        "diárias guardadas em R2, os sinais técnicos e o consenso de "
                        "analistas, e as manchetes das últimas 48 horas."),
        ("De onde", "Yahoo Finance e TradingView para preço e sinal, Banco Central para "
                    "CDI, Selic, IPCA e PTAX, CoinGecko para ativos digitais, Valor "
                    "Econômico e Google News para as manchetes."),
        ("O que garante", "Cada número traz o seu registro de origem: quem forneceu, quando "
                          "foi buscado e se é definitivo ou de sessão em curso."),
    ]),
    ("02", "Inferência", COPPER, [
        ("O que lê", "Tudo o que o primeiro reuniu, junto da exposição de cada cliente por "
                     "classe de ativo."),
        ("O que decide", "O que, de tudo aquilo, importa para estas carteiras hoje — e em "
                         "que ordem. É ele que escreve o resumo do dia e as linhas de "
                         "O que importa hoje, em português."),
        ("O que garante", "A aritmética da exposição fica no código. O modelo nunca escreve "
                          "um número que não esteja nos fatos que recebeu."),
    ]),
    ("03", "Gatilhos", SLATE, [
        ("O que compara", "Cada indicador contra os limiares configurados de cada carteira, "
                          "e cada desvio de alocação contra a política que a governa."),
        ("Como mostra", "Com uma proximidade — 1,0 é estar no limiar — que o portal desenha "
                        "como barra, de modo que dá para ver o que está perto sem ter "
                        "estourado."),
        ("O que garante", "Código puro, sem modelo. Diz quando uma ação é devida; não diz "
                          "qual, e não executa nenhuma."),
    ]),
]
cw3 = (CW - 0.52 * 2) / 3
for i, (num, head, accent, rows) in enumerate(DETAIL):
    x = ML + i * (cw3 + 0.52)
    hairline(s, x, 2.36, cw3, accent, 2.0)
    tb, tf = textbox(s, x, 2.52, cw3, 0.24)
    para(tf, f"agente {num}", font=MONO, size=8, color=accent, tracking=1.0, first=True, caps=True)
    tb, tf = textbox(s, x, 2.78, cw3, 0.34)
    para(tf, head, font=DISPLAY_R, size=18, color=INK, line=1.05, first=True)
    yy = 3.28
    for lab, body in rows:
        tb, tf = textbox(s, x, yy, cw3, 0.20)
        para(tf, lab, font=BODY_R, size=7.5, color=G500, tracking=1.5, first=True, caps=True)
        tb, tf = textbox(s, x, yy + 0.22, cw3, 0.86)
        para(tf, body, font=BODY, size=9, color=INK, line=1.42, first=True)
        yy += 1.12

band(s, "Sem chave de API o produto não para: as manchetes continuam chegando e são "
        "classificadas por regra, e o resumo do dia sai de um redator determinístico, em "
        "português", ML, 6.50, CW, 0.46, fill=G100, color=G600, size=8)


# ══════════════════════════════════════════ 13 · o quarto agente ═════════

s = new(label="Como funciona")
eyebrow(s, "Como funciona  ·  o quarto agente")
title(s, "A carta mensal, em quatro etapas")
lead(s, "Comandado pelo assessor a partir da página do cliente. Escreve a mesma carta que "
        "o grafo do Rivet produz e a deposita na mesma linha, de modo que existe um único "
        "documento mensal por cliente e um único lugar onde ele é lido, aprovado e "
        "publicado.", h=0.46)

FOUR = [
    ("01", "Dados", "A carteira aprovada, a política, os fluxos do mês, os preços de "
                    "abertura e fechamento de cada posição com o câmbio, e a próxima "
                    "reunião já na agenda — cada um com o registro de quem forneceu."),
    ("02", "Análise", "Rentabilidade do mês por Dietz modificado contra a alocação-alvo da "
                      "política precificada como referência; os indicadores e eventos "
                      "mantidos só onde tocam esta carteira; leitura técnica e consenso por "
                      "posição; e as propostas medidas contra a política."),
    ("03", "Redação", "O modelo escreve o título, a saudação e quatro a seis parágrafos a "
                      "partir de um objeto de fatos em que toda cifra já é texto formatado. "
                      "Nunca lhe pedem um número — e uma carta com cifra que os fatos não "
                      "forneceram é recusada e pedida de novo."),
    ("04", "Diagramação", "A carta na página um e o anexo na página dois, o e-mail e a "
                          "página do portal, todos para o R2 com o relatório canônico por "
                          "trás. Uma terceira página faz a etapa falhar."),
]
y = 2.42
for num, head, body in FOUR:
    numbered(s, num, head, body, ML, y, CW, 0.94, head_size=15, body_size=9.4)
    y += 1.04

hairline(s, ML, 6.42, CW, G300, 0.75)
tb, tf = textbox(s, ML, 6.58, CW, 0.50)
p = para(tf, "", first=True)
run(p, "A carta nasce pendente de aprovação. ", font=BODY_R, size=10, color=INK)
run(p, "Só uma carta publicada chega ao cliente, um mês já publicado é recusado em vez de "
       "sobrescrito, e o mesmo agente percorre o book inteiro quando o assessor pede todas "
       "as cartas de uma vez.", font=BODY, size=10, color=INK)


# ════════════════════════════════════════════ 14 · a base técnica ════════

s = new(label="Como funciona")
eyebrow(s, "Como funciona  ·  a base da construção")
title(s, "Claude Code, a API da Anthropic e um Worker")
lead(s, "A construção foi feita em Claude Code, sobre a API da Anthropic. O modelo tem um "
        "trabalho definido e estreito, e o resto é código.", h=0.40)

TECH = [
    ("O que o modelo faz", "Escreve e classifica: a carta, a leitura do dia e o assunto de "
                           "cada manchete."),
    ("O que ele recebe", "Os números já calculados, como texto formatado. Se devolver um que "
                         "não estava nos fatos, a carta é recusada."),
    ("Sem chave", "Um redator determinístico faz o mesmo em português, e as manchetes são "
                  "classificadas por regra. O produto não depende de o modelo estar de pé."),
    ("O Worker", "Cada etapa é um bloco que chama o Worker da Cloudflare, onde a conta é "
                 "feita. Trocar o driver não muda o resultado."),
    ("O Rivet", "Existe para que a sequência e os prompts fiquem legíveis e editáveis: um "
                "orquestrador e nove estágios encadeados por um gate, gerados a partir do "
                "arquivo único de prompts."),
]
y = 2.36
for lab, body in TECH:
    callout(s, lab, body, ML, y, 7.30, 0.76, lab_w=1.72, body_size=9.2)
    y += 0.80

DX, DY, DW = 8.50, 2.36, 3.74
card(s, DX, DY, DW, 2.96, fill=G100, line=G200)
tb, tf = textbox(s, DX + 0.24, DY + 0.22, DW - 0.48, 0.24)
para(tf, "A cadeia", font=BODY_R, size=8, color=G600, tracking=1.6, first=True, caps=True)
CHAIN = [
    ("Rivet", "o grafo: sequência e prompts, legíveis", SLATE),
    ("Worker", "a conta, na Cloudflare", COPPER),
    ("D1 · R2 · KV", "registro, PDFs e cache", CHARCOAL),
]
yy = DY + 0.58
for i, (name, note, col) in enumerate(CHAIN):
    rect(s, DX + 0.30, yy, DW - 0.60, 0.56, fill=WHITE, line=G200)
    rect(s, DX + 0.30, yy, 0.05, 0.56, fill=col)
    tb, tf = textbox(s, DX + 0.50, yy + 0.09, DW - 0.90, 0.20)
    para(tf, name, font=DISPLAY_R, size=11.5, color=INK, line=1.05, first=True)
    tb, tf = textbox(s, DX + 0.50, yy + 0.30, DW - 0.90, 0.20)
    para(tf, note, font=BODY, size=8, color=G600, line=1.2, first=True)
    if i < 2:
        arrow(s, DX + DW / 2, yy + 0.60, DX + DW / 2, yy + 0.80, COPPER, 1.0)
    yy += 0.82
tb, tf = textbox(s, DX + 0.30, DY + 3.10, DW - 0.60, 0.60)
para(tf, "O modelo entra em um estágio só, o da redação, e recebe fatos — nunca a planilha.",
     font=BODY, size=8.8, color=INK, line=1.4, first=True)

source(s, "src/llm/prompts.js — todo prompt em um arquivo; rivet/build-graph.mjs gera o grafo a partir dele.")


# ═══════════════════════════════════════════ 15 · a planta na Cloudflare ══

s = new(label="Como funciona")
eyebrow(s, "Como funciona  ·  a planta")
title(s, "Onde tudo mora, na Cloudflare")
lead(s, "Um Worker à frente, três Workflows duráveis atrás, e três armazéns com papéis "
        "separados: o registro, os arquivos, o cache.", h=0.34)

# ── a fronteira: o que roda na Cloudflare e o que a aciona de fora ─────────
BX, BY, BW, BH = 2.62, 2.14, MR - 2.62, 2.46
card(s, BX, BY, BW, BH, fill=WHITE, line=G300)
rect(s, BX + 0.34, BY - 0.10, 1.30, 0.20, fill=WHITE)          # knockout: the rule stops for the label
tb, tf = textbox(s, BX + 0.40, BY - 0.09, 1.30, 0.20)
para(tf, "Cloudflare", font=BODY_R, size=7.5, color=COPPER, tracking=2.0, first=True, caps=True)

TRIG = [
    ('clock',    "Cron",       "todo dia às 10:00 UTC — 07:00 em São Paulo"),
    ('cursor',   "O assessor", "pede a carta de um cliente, ou o book inteiro"),
    ('terminal', "Job Node",   "09:00 UTC: o macro da XP, o site que barra o Worker"),
]
ty = BY + 0.06
for ic, head, note in TRIG:
    icon(s, ic, ML, ty + 0.01, 0.26, SLATE, 1.0)
    tb, tf = textbox(s, ML + 0.34, ty, 1.30, 0.20)
    para(tf, head, font=DISPLAY_R, size=11, color=INK, line=1.05, first=True)
    tb, tf = textbox(s, ML, ty + 0.28, 1.55, 0.46)
    para(tf, note, font=BODY, size=7.6, color=G600, line=1.34, first=True)
    arrow(s, ML + 1.60, ty + 0.13, BX + 0.22, ty + 0.13, COPPER, 1.0)
    ty += 0.87

# ── o Worker ──────────────────────────────────────────────────────────────
WX, WY, WW, WH = 2.86, 2.42, 2.34, 1.84
card(s, WX, WY, WW, WH, fill=G100, line=G200)
rect(s, WX, WY, WW, 0.05, fill=COPPER)
icon(s, 'chip', WX + 0.18, WY + 0.22, 0.30, COPPER, 1.1)
tb, tf = textbox(s, WX + 0.58, WY + 0.24, WW - 0.72, 0.28)
para(tf, "Worker", font=DISPLAY_R, size=16, color=INK, line=1.05, first=True)
tb, tf = textbox(s, WX + 0.18, WY + 0.60, WW - 0.36, 0.18)
para(tf, "enter-wealth-advisor", font=MONO, size=7.2, color=G500, first=True)
bullets(s, ["o gate, antes da página", "as rotas /api e os ASSETS", "as saídas partem daqui"],
        WX + 0.18, WY + 0.88, WW - 0.36, size=8.2, gap=0.10, indent=0.20, mark=0.065)

# ── os Workflows ──────────────────────────────────────────────────────────
FX, FY, FW = 5.57, WY, 2.70
card(s, FX, FY, FW, WH, fill=G100, line=G200)
rect(s, FX, FY, FW, 0.05, fill=COPPER)
icon(s, 'flow', FX + 0.18, FY + 0.22, 0.30, COPPER, 1.1)
tb, tf = textbox(s, FX + 0.58, FY + 0.24, FW - 0.72, 0.28)
para(tf, "Workflows", font=DISPLAY_R, size=16, color=INK, line=1.05, first=True)
tb, tf = textbox(s, FX + 0.18, FY + 0.60, FW - 0.36, 0.18)
para(tf, "duráveis — a aba pode fechar", font=BODY, size=7.6, color=G500, first=True)
WF = [("Panorama", "3 etapas"), ("Carta mensal", "4 etapas"), ("Cartas em lote", "1 por cliente")]
fy = FY + 0.86
for name, note in WF:
    rect(s, FX + 0.18, fy + 0.085, 0.055, 0.055, fill=COPPER)
    tb, tf = textbox(s, FX + 0.34, fy, 1.50, 0.20)
    para(tf, name, font=DISPLAY_R, size=10.5, color=INK, line=1.05, first=True)
    tb, tf = textbox(s, FX + FW - 1.10, fy + 0.025, 0.92, 0.18)
    para(tf, note, font=MONO, size=7, color=G500, align=PP_ALIGN.RIGHT, first=True)
    hairline(s, FX + 0.18, fy + 0.24, FW - 0.36, G200, 0.75)
    fy += 0.32

# ── os três armazéns ──────────────────────────────────────────────────────
SX, SW = 8.64, 3.36
STORES = [
    ('database', "D1", "DB", "clientes, política, recomendações e aprovações", COPPER),
    ('bucket',   "R2", "REPORTS · MARKET_SERIES", "PDFs e cartas prontas · séries diárias", SLATE),
    ('key',      "KV", "MARKET_CACHE", "cache dos provedores — nunca autoritativo", OLIVE),
]
sy = WY
for ic, name, binding, note, col in STORES:
    card(s, SX, sy, SW, 0.58)
    rect(s, SX, sy, 0.05, 0.58, fill=col)
    icon(s, ic, SX + 0.20, sy + 0.16, 0.26, col, 1.05)
    tb, tf = textbox(s, SX + 0.56, sy + 0.10, 1.00, 0.22)
    para(tf, name, font=DISPLAY_R, size=13, color=INK, line=1.05, first=True)
    tb, tf = textbox(s, SX + 1.50, sy + 0.14, SW - 1.68, 0.18)
    para(tf, binding, font=MONO, size=6.8, color=G500, align=PP_ALIGN.RIGHT, first=True)
    tb, tf = textbox(s, SX + 0.56, sy + 0.33, SW - 0.74, 0.18)
    para(tf, note, font=BODY, size=7.4, color=G600, line=1.2, first=True)
    sy += 0.63

arrow(s, WX + WW + 0.06, WY + 0.92, FX - 0.06, WY + 0.92, COPPER, 1.25)
arrow(s, FX + FW + 0.06, WY + 0.92, SX - 0.06, WY + 0.92, COPPER, 1.25)

# O Worker não fala com os armazéns só através dos Workflows: a página pede, ele lê.
vline(s, WX + WW / 2, WY + WH, 0.10, G400, 0.75)
hairline(s, WX + WW / 2, WY + WH + 0.10, SX + SW / 2 - WX - WW / 2, G400, 0.75)
vline(s, SX + SW / 2, WY + WH, 0.10, G400, 0.75)
rect(s, SX + SW / 2 - 0.03, WY + WH + 0.07, 0.06, 0.06, fill=G400)
tb, tf = textbox(s, WX + WW / 2, WY + WH + 0.16, SX + SW / 2 - WX - WW / 2, 0.18)
para(tf, "e direto, sem agente nenhum: a sessão, a carteira que a página pede, o PDF já arquivado",
     font=BODY, size=7.4, color=G500, align=PP_ALIGN.CENTER, first=True)

# ── o que o Worker chama lá fora ───────────────────────────────────────────
OUTBOUND = [
    ('spark', "API da Anthropic", "escreve a carta e a leitura do dia, e classifica cada manchete", ML),
    ('globe', "Yahoo Finance · Valor · Google News", "preços, séries diárias e as manchetes de 48 horas", 6.30),
]
for ic, head, note, x in OUTBOUND:
    icon(s, ic, x, 4.76, 0.26, SLATE, 1.0)
    tb, tf = textbox(s, x + 0.36, 4.74, 5.0, 0.18)
    para(tf, head, font=BODY_R, size=8.5, color=SLATE, tracking=1.7, first=True, caps=True)
    tb, tf = textbox(s, x + 0.36, 4.92, 5.3, 0.18)
    para(tf, note, font=BODY, size=8, color=G600, first=True)

# ── e, agente por agente, o que cada um toca ───────────────────────────────
TY, RH, HH = 5.20, 0.27, 0.28
CW_ = [1.80, 5.24, 1.10, 1.10, 1.10, 1.10]
CX = [ML]
for w in CW_[:-1]:
    CX.append(CX[-1] + w)

rect(s, ML, TY, CW, HH, fill=CHARCOAL)
for i, lab in enumerate(("Agente", "o que faz, e o que escreve")):
    tb, tf = textbox(s, CX[i] + 0.14, TY, CW_[i] - 0.28, HH, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, lab, font=BODY_R, size=7.5, color=OFFWHITE, tracking=1.2, first=True, caps=True)
for i, (ic, lab) in enumerate((('database', "D1"), ('bucket', "R2"), ('key', "KV"), ('spark', "Modelo"))):
    cx = CX[2 + i] + CW_[2 + i] / 2
    wtxt = len(lab) * (7.5 * 0.56 + 1.2) / 72
    x0 = cx - (0.17 + 0.07 + wtxt) / 2
    icon(s, ic, x0, TY + (HH - 0.17) / 2, 0.17, OFFWHITE, 0.75)
    tb, tf = textbox(s, x0 + 0.24, TY, wtxt + 0.24, HH, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, lab, font=BODY_R, size=7.5, color=OFFWHITE, tracking=1.2, first=True, caps=True)

ROWS = [
    ("01", "Dados", "Cotações, séries diárias e manchetes de 48 horas — cada número com a fonte anotada.",
     "w", "w", "w", "m"),
    ("02", "Inferência", "Lê o que o primeiro reuniu e a exposição do book, e escreve a leitura do dia.",
     "w", "-", "-", "m"),
    ("03", "Gatilhos", "Compara cada indicador e cada desvio com a política que governa a carteira.",
     "w", "-", "-", "-"),
    ("04", "Carta mensal", "Dados, análise, redação, diagramação — a carta escrita e o PDF arquivado.",
     "w", "w", "-", "m"),
    ("", "Cartas em lote", "O mesmo agente, uma etapa durável por cliente do book, e um zip ao final.",
     "w", "w", "-", "m"),
]
ry = TY + HH
for num, name, desc, *marks in ROWS:
    cy = ry + RH / 2
    if num:
        tb, tf = textbox(s, ML + 0.02, ry, 0.34, RH, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, num, font=MONO, size=7.5, color=COPPER, first=True)
    tb, tf = textbox(s, ML + 0.38, ry, CW_[0] - 0.42, RH, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, name, font=DISPLAY_R, size=9.8, color=INK, line=1.1, first=True)
    tb, tf = textbox(s, CX[1] + 0.14, ry, CW_[1] - 0.28, RH, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, desc, font=BODY, size=8.2, color=INK, line=1.2, first=True)
    for i, m in enumerate(marks):
        cx = CX[2 + i] + CW_[2 + i] / 2
        if m == 'w':
            rect(s, cx - 0.045, cy - 0.045, 0.09, 0.09, fill=COPPER)
        elif m == 'r':
            rect(s, cx - 0.045, cy - 0.045, 0.09, 0.09, fill=None, line=G400, line_w=0.75)
        elif m == 'm':
            icon(s, 'spark', cx - 0.075, cy - 0.075, 0.15, COPPER, 0.9)
        else:
            hairline(s, cx - 0.055, cy, 0.11, G300, 1.0)
    hairline(s, ML, ry + RH, CW, G200, 0.75)
    ry += RH

LEGEND = [(7.56, 'w', "escreve"), (8.76, 'r', "só lê"), (9.76, '-', "não toca"), (10.84, 'm', "chama o modelo")]
for x, kind, lab in LEGEND:
    if kind == 'w':
        rect(s, x, 6.975, 0.075, 0.075, fill=COPPER)
    elif kind == 'r':
        rect(s, x, 6.975, 0.075, 0.075, fill=None, line=G400, line_w=0.75)
    elif kind == 'm':
        icon(s, 'spark', x - 0.02, 6.945, 0.13, COPPER, 0.85)
    else:
        hairline(s, x, 7.012, 0.09, G300, 1.0)
    tb, tf = textbox(s, x + 0.16, 6.945, 1.10, 0.18)
    para(tf, lab, font=BODY, size=7.2, color=G600, first=True)

source(s, "wrangler.toml — os bindings desta planta; worker/src/agents.js e letter-agent.js — as etapas.", w=6.9)


# ═══════════════════════════════════════════════ 16 · abertura · IV ══════

s = new(dark=True, right="Capítulo 04")
eyebrow(s, "Capítulo 04", y=2.30, color=COPPER, size=9)
tb, tf = textbox(s, ML, 2.66, 10.6, 1.5)
para(tf, "Next steps", font=DISPLAY_XL, size=42, color=WHITE, line=1.1, first=True)
hairline(s, ML, 3.86, 1.30, COPPER, 1.5)
tb, tf = textbox(s, ML, 4.14, 9.2, 0.8)
para(tf, "Tudo o que foi construído roda sobre dados reais de mercado e dados simulados de "
         "clientes. O caminho até dados reais passa por cinco frentes.",
     font=BODY, size=12.5, color=BFB, line=1.5, first=True)


# ═══════════════════════════════════════════ 17 · real e simulado ════════

s = new(label="Próximos passos")
eyebrow(s, "Próximos passos  ·  onde estamos hoje")
title(s, "Mercado real, clientes simulados")
tb, tf = textbox(s, ML, Y_LEAD, CW, 0.52)
p = para(tf, "", first=True)
p.line_spacing = 1.45
run(p, "A separação não é uma nota de rodapé: o dado simulado carrega ", font=BODY, size=11.5, color=INK)
run(p, "mocked: true", font=MONO, size=10, color=SLATE)
run(p, " no seu registro de origem, mostra o selo SIMULADO na aba de auditoria do assessor "
       "e é nomeado na linha de fontes da própria carta.", font=BODY, size=11.5, color=INK)

COLW = (CW - 0.60) / 2
for i, (head, items, col) in enumerate([
    ("Real, buscado ao vivo", [
        "Todo preço listado — Yahoo Finance e TradingView.",
        "CDI, Selic, IPCA e PTAX — Banco Central.",
        "Leitura técnica e consenso de analistas — TradingView.",
        "As manchetes do dia — Valor Econômico e Google News, cada uma ligada ao seu artigo.",
        "A referência da política, composta a partir de tudo acima.",
    ], SLATE),
    ("Simulado, e rotulado", [
        "As cotas mensais dos fundos brasileiros — não há feed público; a resposta honesta "
        "é o extrato do custodiante.",
        "O histórico de retorno anterior à existência da plataforma.",
        "Os quatro clientes de demonstração além de Albert da Silva.",
    ], COPPER),
]):
    x = ML + i * (COLW + 0.60)
    hairline(s, x, 2.46, COLW, col, 2.0)
    tb, tf = textbox(s, x, 2.62, COLW, 0.30)
    para(tf, head, font=DISPLAY_R, size=15, color=INK, line=1.1, first=True)
    bullets(s, items, x, 3.06, COLW, size=9.6, color=col, gap=0.22)

hairline(s, ML, 5.40, CW, G300, 0.75)
tb, tf = textbox(s, ML, 5.58, CW, 0.30)
para(tf, "Também é real, e vem dos arquivos do caso", font=BODY_R, size=8.5, color=G600,
     tracking=1.7, first=True, caps=True)
tb, tf = textbox(s, ML, 5.90, CW, 0.70)
para(tf, "A identidade de Albert, o assessor e o perfil de risco; o extrato de posição da "
         "XP de maio de 2025, reproduzido como snapshot v1; e as projeções macro da XP de "
         "fevereiro de 2025 — usadas como o que são, um documento datado, e não como se "
         "fossem a leitura de hoje.", font=BODY, size=9.6, color=INK, line=1.45, first=True)


# ═════════════════════════════════════ 18 · frentes 1 a 3 ═══════════════

s = new(label="Próximos passos")
eyebrow(s, "Próximos passos  ·  frentes 1 a 3")
title(s, "Acesso, nuvem e segurança do dado")
lead(s, "Cinco frentes separam o que está de pé do primeiro cliente real. As três "
        "primeiras são de acesso, de contrato e de infraestrutura.", h=0.40)

y = 2.30
y = block(s, "01", "Acesso", [
    "A senha compartilhada do portal sai e entra o SSO corporativo da XP com MFA pelo "
    "Cloudflare Access — que o código já prevê.",
    "A regra de que cada assessor enxerga apenas o próprio livro é verificada hoje num "
    "único ponto do sistema, e deve continuar assim.",
], ML, y, CW)
y = block(s, "02", "Compliance com a nuvem", [
    "A Cloudflare precisa ser contratada como a CMN 4.893 exige: comunicação ao Banco "
    "Central, os requisitos adicionais para processamento no exterior, localização dos "
    "dados, direito de auditoria e plano de saída em contrato.",
    "O mesmo vale para a Anthropic, com acordo de não-retenção dos prompts.",
], ML, y, CW)
y = block(s, "03", "Segurança do dado financeiro", [
    "Criptografia com chaves geridas pela instituição, segregação por assessor e rotação "
    "de segredos.",
    "Trilha de auditoria exportável, e um teste de intrusão independente antes do primeiro "
    "cliente real.",
], ML, y, CW, rule=False)

source(s, "CMN 4.893/2021 — política de segurança cibernética e contratação de processamento e armazenamento de dados em nuvem.")


# ═════════════════════════════════════ 19 · frentes 4 e 5 ═══════════════

s = new(label="Próximos passos")
eyebrow(s, "Próximos passos  ·  frentes 4 e 5")
title(s, "Sigilo, e governança do que o modelo escreve")
lead(s, "As duas últimas são sobre dado pessoal e sobre o que uma pessoa precisa ter lido "
        "antes de a carta sair.", h=0.40)

y = 2.30
y = block(s, "04", "Sigilo", [
    "O sigilo bancário da LC 105 e a LGPD pedem minimização. O modelo já recebe só o "
    "objeto de fatos, e não a base.",
    "O próximo passo é pseudonimizar nome e documento no prompt, definir retenção para "
    "cartas e logs, e registrar quem viu o quê.",
], ML, y, CW)
y = block(s, "05", "Governança do que o modelo escreve", [
    "Manter a aprovação do assessor como obrigatória: nenhuma carta chega ao cliente sem "
    "ter sido lida por uma pessoa.",
    "Versionar os prompts — já existe — e medir a qualidade das cartas contra um conjunto "
    "fixo de casos a cada mudança.",
    "Assim, uma troca de modelo nunca chega ao cliente sem ter sido avaliada antes.",
], ML, y, CW, rule=False)

hairline(s, ML, 6.26, CW, G300, 0.75)
tb, tf = textbox(s, ML, 6.44, CW, 0.40)
p = para(tf, "", first=True)
run(p, "O princípio que atravessa as cinco frentes:  ", font=BODY, size=12, color=INK)
run(p, "o que é verificável fica em código, e o que é escrito passa por uma pessoa.",
    font=DISPLAY_R, size=13, color=COPPER)

source(s, "LC 105/2001 — sigilo das operações de instituições financeiras; Lei 13.709/2018 — LGPD.")


# ═════════════════════════════════════════════ 20 · fechamento ═══════════

s = new(dark=True, right="Obrigado")
eyebrow(s, "Para ver ao vivo", y=1.86, color=COPPER, size=9)
tb, tf = textbox(s, ML, 2.22, 10.6, 1.3)
para(tf, "Uma carta por cliente, por mês —", font=DISPLAY_XL, size=34, color=WHITE, line=1.14, first=True)
para(tf, "lida por uma pessoa antes de sair.", font=DISPLAY_XL, size=34, color=WHITE, line=1.14)
hairline(s, ML, 4.02, 1.30, COPPER, 1.5)

LINKS = [
    ("Portal", "enter-wealth-advisor.joaocreste-8da.workers.dev"),
    ("Assessor", "antonio.bicudo@xpi.com.br"),
    ("Cliente", "albert.dasilva@exemplo.com.br"),
]
x = ML
for lab, val in LINKS:
    tb, tf = textbox(s, x, 4.42, 4.2, 0.20)
    para(tf, lab, font=BODY, size=7.5, color=SAGE, tracking=1.9, first=True, caps=True)
    tb, tf = textbox(s, x, 4.66, 4.2, 0.30)
    para(tf, val, font=MONO, size=9, color=DDD, line=1.3, first=True)
    x += 4.2

tb, tf = textbox(s, ML, 5.40, 9.0, 0.40)
para(tf, "O portal está atrás de uma senha de acesso; as credenciais de demonstração vão "
         "com o material.", font=BODY, size=9.5, color=G500, line=1.4, first=True)

prs.save(OUT)
print(f"{OUT}  ·  {page['n']} slides")
